"""
Rewrite IMU_外骨骼輔助(1).docx — clean, compatible with all Word editors.
Content preserved from original; formulas corrected per 2025 revision.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

IMG_DIR = r"C:\Users\long\Desktop\exo\doc_imgs"

def add_image(doc, filename, width_inches=5.5, caption=None, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run()
    run.add_picture(f"{IMG_DIR}\\{filename}", width=Inches(width_inches))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after  = Pt(12)
        cr = cp.add_run(caption)
        cr.font.size = Pt(9)
        cr.font.name = "Arial"
        cr.font.italic = True
        cr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# Page: A4, 1-inch margins
sec = doc.sections[0]
sec.page_width  = Inches(8.27)
sec.page_height = Inches(11.69)
for attr in ("left_margin","right_margin","top_margin","bottom_margin"):
    setattr(sec, attr, Inches(1.0))

# ─── Helpers ──────────────────────────────────────────────────────────────────
BLUE_DARK  = RGBColor(0x1F, 0x38, 0x82)
BLUE_MED   = RGBColor(0x2E, 0x75, 0xB6)
BLUE_LIGHT = RGBColor(0x26, 0x51, 0x82)
RED        = RGBColor(0xC0, 0x00, 0x00)
GREEN      = RGBColor(0x37, 0x86, 0x30)
GRAY       = RGBColor(0x59, 0x59, 0x59)

def para(text="", bold=False, size=11, color=None, align=None, space_before=0, space_after=6, indent_cm=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    if indent_cm:
        p.paragraph_format.left_indent = Cm(indent_cm)
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Arial"
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p

def heading(text, level=1):
    """Simple heading without using heading styles (for compatibility)."""
    sizes = {1: 15, 2: 13, 3: 11}
    colors = {1: BLUE_MED, 2: BLUE_LIGHT, 3: BLUE_DARK}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after  = Pt(6)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(sizes[level])
    r.font.color.rgb = colors[level]
    return p

def formula(text, note=""):
    """Box-style formula paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Cm(1.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    r = p.add_run(text)
    r.font.name = "Courier New"
    r.font.size = Pt(11)
    r.font.color.rgb = BLUE_DARK
    if note:
        rn = p.add_run(f"    ← {note}")
        rn.font.name = "Arial"
        rn.font.size = Pt(9)
        rn.font.color.rgb = GRAY
    return p

def bullet(text, indent=1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Cm(indent * 0.8)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(3)
    r = p.add_run("• " + text)
    r.font.name = "Arial"
    r.font.size = Pt(11)
    return p

def set_cell(cell, text, bold=False, center=False, fill_hex=None, size=10, color=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if fill_hex:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill_hex)
        tcPr.append(shd)

def make_table(headers, rows, col_widths_cm, alt_fill="EBF3FB", header_fill="2E75B6"):
    n = len(headers)
    tbl = doc.add_table(rows=1+len(rows), cols=n)
    tbl.style = "Table Grid"
    # header row
    for i, h in enumerate(headers):
        set_cell(tbl.rows[0].cells[i], h, bold=True, center=True,
                 fill_hex=header_fill, size=10,
                 color=RGBColor(0xFF, 0xFF, 0xFF))
    # data rows
    for ri, row in enumerate(rows):
        fill = alt_fill if ri % 2 == 0 else None
        for ci, val in enumerate(row):
            set_cell(tbl.rows[ri+1].cells[ci], str(val), fill_hex=fill,
                     center=(ci != 1 and len(row) > 2))
    # set column widths
    for row in tbl.rows:
        for ci, w in enumerate(col_widths_cm):
            row.cells[ci].width = Cm(w)
    return tbl

# ═══════════════════════════════════════════════════════════════════════════════
# TITLE
# ═══════════════════════════════════════════════════════════════════════════════
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_before = Pt(0)
p_title.paragraph_format.space_after  = Pt(4)
r = p_title.add_run("IMU 外骨骼輔助 — 人機介面開發需求文件")
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(18)
r.font.color.rgb = BLUE_DARK

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_after = Pt(16)
r2 = p_sub.add_run("外骨骼腰椎壓力評估系統  |  WIT WT901BLE IMU 版本  |  2025 修訂")
r2.font.name = "Arial"
r2.font.size = Pt(10)
r2.font.color.rgb = GRAY

para("利用 IMU 讀取動態姿勢並估算軀幹前傾角（Trunk Flexion Angle），評估穿戴外骨骼後對腰椎壓力的降低效果。")

# ═══════════════════════════════════════════════════════════════════════════════
# 一、資料庫欄位
# ═══════════════════════════════════════════════════════════════════════════════
heading("一、資料庫輸入欄位定義", level=1)

make_table(
    headers=["欄位名稱", "中文說明", "單位", "備註"],
    col_widths_cm=[4.0, 4.2, 2.0, 5.5],
    rows=[
        ("record_id",         "紀錄編號",       "—",     "自動產生，格式 001–999"),
        ("subject_name",      "姓名",           "—",     "中英文均可"),
        ("gender",            "性別",           "—",     "男 / 女"),
        ("age",               "年齡",           "歲",    ""),
        ("height_cm",         "身高",           "cm",    ""),
        ("weight_kg",         "體重",           "kg",    "計算 W_upper 的基礎"),
        ("spring_k",          "彈簧係數",       "Nm/deg","被動式外骨骼彈簧剛度"),
        ("task_type",         "任務類型",       "—",     "擺位 / 移床"),
        ("condition",         "穿戴條件",       "—",     "bare（未穿）/ exo（穿戴）"),
        ("timestamp",         "時間戳記",       "—",     "自動代入電腦時間"),
        ("theta_deg",         "軀幹前傾角",     "°",     "θ = θ1 − θ2，校正後"),
        ("M_lumbar",          "腰椎平衡力矩",   "N·m",   ""),
        ("M_exo",             "外骨骼輔助力矩", "N·m",   ""),
        ("M_mus_exo",         "背肌剩餘力矩",   "N·m",   ""),
        ("F_c_bare",          "未穿戴腰椎壓迫力","N",    ""),
        ("F_c_exo",           "穿戴腰椎壓迫力", "N",     ""),
        ("reduction_percent", "外骨骼有效率",   "%",     "壓力降低百分比"),
        ("load_index",        "負載指數",       "—",     "F_c / 3400"),
        ("risk_flag",         "風險等級",       "—",     "normal / warning / high"),
    ]
)

# ═══════════════════════════════════════════════════════════════════════════════
# 二、按鈕功能
# ═══════════════════════════════════════════════════════════════════════════════
heading("二、介面按鈕功能說明", level=1)

make_table(
    headers=["按鈕", "功能說明"],
    col_widths_cm=[3.5, 12.0],
    rows=[
        ("IMU 校正",  "受測者站直後按下，將目前 IMU 姿態設定為 0°（消除初始偏移）"),
        ("開始測試",  "開始紀錄，每 3 秒自動儲存一次當前數據至資料庫"),
        ("停止測試",  "停止自動紀錄，保留已存數據"),
        ("儲存紀錄",  "立即手動儲存目前時間點的數據"),
        ("匯出報表",  "將資料庫匯出為 Excel 彙整分析表"),
    ]
)

# ═══════════════════════════════════════════════════════════════════════════════
# 三、記錄條件
# ═══════════════════════════════════════════════════════════════════════════════
heading("三、記錄條件與 IMU 角度計算", level=1)

add_image(doc, "fig3_imu.png", width_inches=3.2,
          caption="圖 1  IMU 配戴位置示意圖（藍色 = 軀幹 IMU1，綠色 = 骨盆 IMU2）")

heading("3.1  軀幹前傾角計算", level=2)
para("若使用兩顆 IMU（軀幹 + 骨盆），則：")
formula("θ = θ1 − θ2")

make_table(
    headers=["符號", "說明"],
    col_widths_cm=[3.0, 12.5],
    rows=[
        ("θ1", "軀幹 IMU 量測角度（貼附於上背部）"),
        ("θ2", "骨盆 / 腰部 IMU 量測角度（貼附於骨盆後側）"),
        ("θ",  "校正後軀幹前傾角（IMU 校正後歸零，動作時為正值）"),
    ]
)

para("")
heading("3.2  L5-S1 解剖位置說明", level=2)
make_table(
    headers=["名稱", "位置說明"],
    col_widths_cm=[3.5, 12.0],
    rows=[
        ("L5（第五腰椎）", "腰椎最下方一節椎骨，位於骨盆正上方，是人體承受負重最大的腰椎關節之一"),
        ("S1（第一薦椎）", "薦骨（Sacrum）最上端，位於兩側髂骨之間"),
        ("L5-S1 椎間盤",   "本系統評估腰椎壓迫力（F_c）的計算基準位置"),
    ]
)

# ═══════════════════════════════════════════════════════════════════════════════
# 四、測試步驟
# ═══════════════════════════════════════════════════════════════════════════════
heading("四、測試步驟", level=1)

steps = [
    ("STEP 1", "受測者站直，按下「IMU 校正」鈕，將當前姿態設為 0°。"),
    ("STEP 2", "受測者未穿外骨骼，配戴 IMU，執行「擺位」或「移床」動作，按開始測試後系統每 3 秒自動記錄一次。"),
    ("STEP 3", "受測者同時穿戴外骨骼與 IMU，重複執行相同動作並記錄（condition = exo）。"),
    ("STEP 4", "可連續測試，系統自動統計腰椎最大壓迫力，完成後按「匯出報表」匯出 Excel 分析結果。"),
]
make_table(
    headers=["步驟", "說明"],
    col_widths_cm=[2.0, 13.5],
    rows=steps
)

# ═══════════════════════════════════════════════════════════════════════════════
# 五、人機介面呈現數據
# ═══════════════════════════════════════════════════════════════════════════════
heading("五、人機介面呈現數據", level=1)

make_table(
    headers=["顯示項目", "單位", "說明"],
    col_widths_cm=[4.5, 2.0, 9.0],
    rows=[
        ("軀幹前傾角 θ",     "°",    "Trunk Flexion Angle，IMU 即時量測"),
        ("F_c_bare",         "N",    "未穿戴外骨骼估算腰椎壓迫力"),
        ("F_c_exo",          "N",    "穿戴外骨骼估算腰椎壓迫力"),
        ("M_lumbar",         "N·m",  "腰椎平衡力矩"),
        ("M_exo",            "N·m",  "外骨骼輔助力矩"),
        ("M_mus_exo",        "N·m",  "穿戴後背肌仍需承擔的力矩"),
        ("外骨骼有效率",     "%",    "壓力降低率 = (F_c_bare − F_c_exo) / F_c_bare × 100"),
        ("動作總時間",       "s",    "每次擺位 / 移床的動作時間"),
        ("腰椎最大壓迫力",   "N",    "該次測試最大 F_c 值"),
        ("負載指數 LI",      "—",    "F_c / 3400，無因次風險指標"),
    ]
)

para("")
para("負載指數風險判讀：", bold=True)
make_table(
    headers=["Load Index (LI)", "風險等級", "說明"],
    col_widths_cm=[4.0, 4.0, 7.5],
    alt_fill=None,
    rows=[
        ("LI < 0.75",       "低風險（normal）",   "F_c < 2550 N，操作安全"),
        ("0.75 ≤ LI ≤ 1.0", "中等風險（warning）","F_c = 2550–3400 N，需留意"),
        ("LI > 1.0",        "高風險（high）",      "F_c > 3400 N，超過 NIOSH 臨界值"),
    ]
)

# ═══════════════════════════════════════════════════════════════════════════════
# 六、生物力學模型（修正版）
# ═══════════════════════════════════════════════════════════════════════════════
heading("六、腰椎壓力生物力學模型（L5-S1）", level=1)

add_image(doc, "fig1_model.png", width_inches=5.0,
          caption="圖 2  腰椎壓力生物力學模型示意圖（準靜態單肌肉等效模型）")

para(
    "腰椎所受的總壓力（Compression Force, F_c）主要來自兩部分：\n"
    "（1）上半身重力在脊椎軸向的分力；（2）豎脊肌收縮維持平衡所產生的壓力。\n"
    "本模型為準靜態、單肌肉等效模型，適用於緩速動作（擺位、移床）。"
)

# 6.1 Parameters
heading("6.1  參數定義", level=2)
make_table(
    headers=["符號", "說明", "單位", "預設值"],
    col_widths_cm=[2.5, 6.5, 2.0, 4.5],
    rows=[
        ("θ",        "校正後軀幹前傾角",                "°",     "IMU 即時量測"),
        ("W_body",   "受測者全身體重",                  "kg",    "介面輸入"),
        ("W_upper",  "腰部以上有效質量（= 0.6 × W_body）","kg",  "0.6 × W_body"),
        ("g",        "重力加速度",                      "m/s²",  "9.81"),
        ("d_COM",    "上半身質心到 L5-S1 距離",         "m",     "0.30（範圍 0.25–0.35）"),
        ("d_mus",    "豎脊肌到脊椎中心有效力臂",        "m",     "0.05（解剖學常數）"),
        ("k",        "被動式外骨骼彈簧係數",            "Nm/deg","使用者輸入"),
        ("NIOSH",    "腰椎傷害臨界值",                  "N",     "3400"),
    ]
)

# 6.2 Bare model
heading("6.2  未穿戴外骨骼基準模型（Baseline）", level=2)

para("Step 1 — 上半身有效質量：")
formula("W_upper = 0.6 × W_body", "[kg]")

para("Step 2 — 腰椎平衡力矩：")
formula("M_lumbar = W_upper × g × d_COM × sin(θ)", "[N·m]")
para("θ 越大，sin(θ) 越大 → M_lumbar 非線性飆升（主要風險驅動因子）", indent_cm=1.5, size=10)

para("Step 3 — 豎脊肌收縮力（由力矩平衡推導）：")
formula("F_mus_bare = M_lumbar / d_mus", "[N]")

para("Step 4 — 未穿戴腰椎壓迫力（垂直重力分量 + 肌肉力）：")
formula("F_c_bare = W_upper × g × cos(θ) + F_mus_bare", "[N]")

# 6.3 Exo model
heading("6.3  穿戴外骨骼後的壓力模型", level=2)

para("Step 5 — 被動式彈簧外骨骼輔助力矩（校正後直立位 θ=0° 時無預壓）：")
formula("M_exo = k × θ", "[N·m]  （k 單位：Nm/deg）")

para("Step 6 — 外骨骼分擔後背肌仍需承擔的剩餘力矩：")
formula("M_mus_exo = max(0, M_lumbar − M_exo)", "[N·m]")
para("max(0, …) 確保外骨骼不對背肌產生「負力矩」（過度輔助時設為 0）", indent_cm=1.5, size=10)

para("Step 7 — 穿戴外骨骼後腰椎壓迫力：")
formula("F_c_exo = W_upper × g × cos(θ) + M_mus_exo / d_mus", "[N]")

# 6.4 Reduction
heading("6.4  外骨骼減壓率", level=2)
formula("Reduction% = (F_c_bare − F_c_exo) / F_c_bare × 100", "[%]  分母為基準值 F_c_bare")
para("設計良好的腰部外骨骼，在 θ > 45° 時通常能降低腰椎壓力 20%–40%。", indent_cm=1.5, size=10)

# 6.5 LI
heading("6.5  負載指數", level=2)
formula("LI = F_c / 3400")
para("以 NIOSH 3400 N 臨界值為基準，F_c 取當前條件（bare 或 exo）的壓迫力。", indent_cm=1.5, size=10)

# ═══════════════════════════════════════════════════════════════════════════════
# 七、數值驗證範例
# ═══════════════════════════════════════════════════════════════════════════════
add_image(doc, "fig2_flow.png", width_inches=5.5,
          caption="圖 3  計算流程圖（對應 calculations.py 實作）")

heading("七、數值驗證範例", level=1)
para("以受測者體重 65 kg、θ = 45°、k = 0.5 Nm/deg 為例：")

make_table(
    headers=["步驟", "計算式", "結果"],
    col_widths_cm=[3.0, 9.0, 3.5],
    rows=[
        ("W_upper",    "0.6 × 65",                                         "39 kg"),
        ("M_lumbar",   "39 × 9.81 × 0.30 × sin(45°)",                     "≈ 81.2 N·m"),
        ("M_exo",      "0.5 × 45",                                         "= 22.5 N·m"),
        ("M_mus_exo",  "81.2 − 22.5",                                      "= 58.7 N·m"),
        ("F_mus_bare", "81.2 / 0.05",                                      "= 1624 N"),
        ("F_c_bare",   "39×9.81×cos(45°) + 1624",                         "≈ 1894 N"),
        ("F_mus_exo",  "58.7 / 0.05",                                      "= 1174 N"),
        ("F_c_exo",    "39×9.81×cos(45°) + 1174",                         "≈ 1444 N"),
        ("Reduction%", "(1894 − 1444) / 1894 × 100",                      "≈ 23.8%"),
        ("LI (bare)",  "1894 / 3400",                                       "≈ 0.557 → 低風險"),
        ("LI (exo)",   "1444 / 3400",                                       "≈ 0.425 → 低風險"),
    ]
)

# ═══════════════════════════════════════════════════════════════════════════════
# 八、模型假設與限制
# ═══════════════════════════════════════════════════════════════════════════════
heading("八、模型假設與適用限制", level=1)

limitations = [
    ("準靜態假設",
     "本模型忽略動態加速度（慣性力），假設動作緩慢且近似靜態平衡。"
     "實際護理搬運涉及加速，真實椎間盤壓力會高於模型估算值。"),
    ("單肌肉等效模型",
     "使用單一等效豎脊肌代表所有腰背伸肌群，為工程近似。"
     "在評估「相對減壓效果」時仍具參考價值。"),
    ("上半身比例固定",
     "W_upper = 0.6 × W_body 為固定比例，文獻範圍為 0.55–0.67。"),
    ("外骨骼線性彈簧模型",
     "M_exo = k × θ 為線性近似，適用於低速動作。"
     "氣壓棒等非線性元件，應改用廠商提供的 Lookup Table 或氣體力學公式。"),
    ("未計入手持物重量",
     "目前 load_kg = 0，若需評估搬運情境，"
     "應將搬運物重量加入：W_upper = 0.6 × W_body + W_load。"),
]

for title_text, body_text in limitations:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(3)
    tr = p.add_run(f"▶  {title_text}：")
    tr.bold = True
    tr.font.name = "Arial"
    tr.font.size = Pt(11)
    tr.font.color.rgb = BLUE_MED
    br = p.add_run(body_text)
    br.font.name = "Arial"
    br.font.size = Pt(10.5)

# Footer
para("")
p_f = doc.add_paragraph()
p_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
rf = p_f.add_run("參考文獻：McGill (2007) | Winter (2009) | NIOSH (1994) | 本文件公式已於 2025 年審查修正")
rf.font.size = Pt(9)
rf.font.name = "Arial"
rf.font.color.rgb = GRAY

doc.save(r"C:\Users\long\Desktop\exo\IMU_exo_revised.docx")
print("Done.")
