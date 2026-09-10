import os
import threading
import time
import zipfile
from collections import defaultdict
from copy import copy
from datetime import datetime
import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from calculations import LEAN_BACK_DEG

DB_PATH = os.path.join(os.path.dirname(__file__), "exo_records.xlsx")
SHEET = "Records"

COLUMNS = [
    "record_id", "subject_name", "gender", "age", "height_cm", "weight_kg",
    "patient_kg", "spring_k", "task_type", "condition", "timestamp",
    "theta_deg", "M_lumbar", "M_exo", "M_mus_exo",
    "F_c_bare", "F_c_exo", "reduction_percent", "load_index", "risk_flag",
]

COLUMN_LABELS = {
    "record_id": "紀錄編號",
    "subject_name": "姓名",
    "gender": "性別",
    "age": "年齡",
    "height_cm": "身高(cm)",
    "weight_kg": "體重(kg)",
    "patient_kg": "老人重量(kg)",
    "spring_k": "彈簧係數",
    "task_type": "任務類型",
    "condition": "條件",
    "timestamp": "時間戳記",
    "theta_deg": "軀幹前傾角(°)",       # signed: negative = leaning back
    "M_lumbar": "腰椎力矩(N·m)",
    "M_exo": "外骨骼力矩(N·m)",
    "M_mus_exo": "背肌剩餘力矩(N·m)",
    "F_c_bare": "未穿戴壓迫力(N)",
    "F_c_exo": "穿戴壓迫力(N)",
    "reduction_percent": "減壓率(%)",
    "load_index": "負載指數",
    "risk_flag": "風險等級",
}

HEADER_FILL = PatternFill("solid", fgColor="2E75B6")
WARNING_FILL = PatternFill("solid", fgColor="FFD966")
HIGH_FILL = PatternFill("solid", fgColor="FF0000")
NORMAL_FILL = PatternFill("solid", fgColor="70AD47")

thin = Side(style="thin", color="CCCCCC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
LEAN_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")

# Header of sheets written by the short-lived version that stored the angle twice
# (signed 量測角度 + clamped 計算用角度); folded back into one column on next write.
_TWO_ANGLE_HEADER = [label for c in COLUMNS
                     for label in (["量測角度(°)", "計算用角度(°)"] if c == "theta_deg"
                                   else [COLUMN_LABELS[c]])]


# ── Safe file access ─────────────────────────────────────────────────────────
# Every browser tab is a session in this one process, and a running test reads
# the workbook on each refresh while writing it every 3 s. openpyxl rewrites the
# file in place, so a read (another tab, Excel) landing mid-write used to see a
# half-written zip -> BadZipFile. Hence:
#   * one lock serialises all reads/writes across sessions;
#   * saves go to a temp file swapped in with os.replace(), so the file on disk
#     is always a complete workbook, even for readers outside this process;
#   * brief Windows file locks (Excel, antivirus, Explorer preview) are retried.
_db_lock = threading.RLock()
_RETRY_DELAYS_S = (0.1, 0.2, 0.4)
_UNREADABLE = (zipfile.BadZipFile, KeyError, EOFError)   # truncated / not an xlsx


class DatabaseBusyError(Exception):
    """The workbook stayed locked by another program (typically Excel)."""


class DatabaseCorruptError(Exception):
    """The workbook file exists but is not a readable xlsx."""


def _busy_msg() -> str:
    name = os.path.basename(DB_PATH)
    folder = os.path.dirname(DB_PATH)
    # Excel keeps an owner file "~$<name>" next to a workbook it has open
    if any(os.path.exists(os.path.join(folder, n)) for n in ("~$" + name, "~$" + name[2:])):
        return f"Excel 正開著 {name}，關閉 Excel 後會自動寫入"
    return f"{name} 暫時被其他程式鎖住"


def _with_retry(fn):
    """Call fn; retry a few times on PermissionError, then raise DatabaseBusyError."""
    for delay in _RETRY_DELAYS_S:
        try:
            return fn()
        except PermissionError:
            time.sleep(delay)
    try:
        return fn()
    except PermissionError as e:
        raise DatabaseBusyError(_busy_msg()) from e


def _load_wb():
    """Open the workbook, or None if it doesn't exist yet. Brief locks and
    half-written files are retried before raising Busy / Corrupt."""
    for delay in (*_RETRY_DELAYS_S, None):
        if not os.path.exists(DB_PATH):
            return None
        try:
            return openpyxl.load_workbook(DB_PATH)
        except PermissionError as e:
            if delay is None:
                raise DatabaseBusyError(_busy_msg()) from e
        except _UNREADABLE as e:
            if delay is None:
                raise DatabaseCorruptError(
                    f"{os.path.basename(DB_PATH)} 不是有效的 Excel 檔（可能在寫入途中被中斷）") from e
        time.sleep(delay)


def _save_wb(wb):
    """Save to a temp file, then atomically swap it in for DB_PATH."""
    tmp = os.path.splitext(DB_PATH)[0] + ".tmp.xlsx"
    _with_retry(lambda: wb.save(tmp))
    try:
        _with_retry(lambda: os.replace(tmp, DB_PATH))
    except DatabaseBusyError:
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise


def _new_wb():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    _write_header(ws)
    return wb


def _open_for_write():
    """Return (workbook, backup). An unreadable existing file is moved aside to
    exo_records_corrupt_<time>.xlsx so recording can continue; backup is that
    file's name (None if nothing was moved)."""
    backup = None
    try:
        wb = _load_wb()
    except DatabaseCorruptError:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = os.path.splitext(DB_PATH)[0] + f"_corrupt_{stamp}.xlsx"
        _with_retry(lambda: os.replace(DB_PATH, backup_path))
        backup = os.path.basename(backup_path)
        wb = None
    if wb is None:
        wb = _new_wb()
    if SHEET not in wb.sheetnames:      # e.g. the sheet was renamed in Excel
        _write_header(wb.create_sheet(SHEET, 0))
    _upgrade_sheet(wb[SHEET])
    _ensure_lean_highlight(wb[SHEET])
    return wb, backup


def _write_header(ws):
    for col_idx, col in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=COLUMN_LABELS.get(col, col))
        cell.font = Font(bold=True, color="FFFFFF", name="Arial")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def _upgrade_sheet(ws):
    """Fold the two-angle layout back into the single signed 軀幹前傾角 column:
    keep 量測角度, or 計算用角度 on rows written before 量測角度 existed."""
    if [c.value for c in ws[1]][:len(_TWO_ANGLE_HEADER)] != _TWO_ANGLE_HEADER:
        return
    col = COLUMNS.index("theta_deg") + 1          # 量測角度; 計算用角度 is col + 1
    for row in range(2, ws.max_row + 1):
        measured, clamped = ws.cell(row, col), ws.cell(row, col + 1)
        if measured.value is None:
            measured.value = clamped.value
            measured._style = copy(clamped._style)
    ws.delete_cols(col + 1)
    _write_header(ws)


def _ensure_lean_highlight(ws):
    """Light-blue 軀幹前傾角 cells below LEAN_BACK_DEG (leaning back). An Excel
    conditional format rather than a baked-in fill, so it follows edits made in Excel."""
    col = get_column_letter(COLUMNS.index("theta_deg") + 1)
    rng = f"{col}2:{col}1048576"
    threshold = f"{LEAN_BACK_DEG:g}"
    for cf in ws.conditional_formatting:
        if str(cf.sqref) == rng:
            for rule in cf.rules:
                rule.formula = [threshold]   # follow a changed LEAN_BACK_DEG
            return
    ws.conditional_formatting.add(
        rng, CellIsRule(operator="lessThan", formula=[threshold], fill=LEAN_FILL))


def _next_record_id(ws) -> str:
    max_row = ws.max_row
    if max_row <= 1:
        return "001"
    last_id = ws.cell(row=max_row, column=1).value
    try:
        num = int(str(last_id)) + 1
    except (TypeError, ValueError):
        num = max_row
    return f"{num:03d}"


def _auto_col_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 30)


def _rows_as_dicts(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []
    headers = list(rows[0])
    return [dict(zip(headers, row)) for row in rows[1:]]


def _append_row(ws, subject: dict, metrics: dict) -> str:
    if ws.max_row == 1 and ws.cell(1, 1).value is None:
        _write_header(ws)

    record_id = _next_record_id(ws)

    row_data = {
        "record_id": record_id,
        "subject_name": subject.get("subject_name", ""),
        "gender": subject.get("gender", ""),
        "age": subject.get("age", ""),
        "height_cm": subject.get("height_cm", ""),
        "weight_kg": subject.get("weight_kg", ""),
        "patient_kg": subject.get("patient_kg", ""),
        "spring_k": subject.get("spring_k", ""),
        "task_type": subject.get("task_type", ""),
        "condition": subject.get("condition", ""),
        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        **metrics,   # a "timestamp" taken at measurement time overrides the one above
    }

    row_idx = ws.max_row + 1
    risk = metrics.get("risk_flag", "normal")

    for col_idx, col in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=row_data.get(col, ""))
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if col == "risk_flag":
            if risk == "high":
                cell.fill = HIGH_FILL
                cell.font = Font(color="FFFFFF", name="Arial")
            elif risk == "warning":
                cell.fill = WARNING_FILL
                cell.font = Font(name="Arial")
            else:
                cell.fill = NORMAL_FILL
                cell.font = Font(color="FFFFFF", name="Arial")
        else:
            cell.font = Font(name="Arial")

    return record_id


def save_records(rows: list[tuple[dict, dict]]) -> tuple[list[str], str | None]:
    """Append [(subject, metrics), ...] in one atomic write.

    Returns (record_ids, backup): backup is the name an unreadable old workbook
    was moved to, else None. Raises DatabaseBusyError if the file stays locked;
    nothing is written then, so the caller can simply retry the same rows."""
    with _db_lock:
        wb, backup = _open_for_write()
        ws = wb[SHEET]
        ids = [_append_row(ws, subject, metrics) for subject, metrics in rows]
        _auto_col_width(ws)
        _save_wb(wb)
    return ids, backup


def save_record(subject: dict, metrics: dict) -> str:
    """Append one row; returns its record id."""
    return save_records([(subject, metrics)])[0][0]


def load_records() -> list[dict]:
    """All rows as dicts keyed by header label. May raise DatabaseBusyError /
    DatabaseCorruptError."""
    with _db_lock:
        wb = _load_wb()
    if wb is None or SHEET not in wb.sheetnames:
        return []
    return _rows_as_dicts(wb[SHEET])


def export_summary(subject_name: str = None) -> str:
    """Export a summary sheet grouped by subject. Returns path."""
    with _db_lock:
        wb, _ = _open_for_write()
        records = _rows_as_dicts(wb[SHEET])

        if subject_name:
            records = [r for r in records if str(r.get("姓名", "")) == subject_name]

        sheet_name = "Summary"
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws_sum = wb.create_sheet(sheet_name)

        headers = ["姓名", "條件", "任務類型", "最大壓迫力(N)", "平均減壓率(%)", "高風險次數", "負載指數(最大)"]
        for ci, h in enumerate(headers, 1):
            cell = ws_sum.cell(1, ci, h)
            cell.font = Font(bold=True, color="FFFFFF", name="Arial")
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = BORDER

        groups = defaultdict(list)
        for r in records:
            key = (r.get("姓名", ""), r.get("條件", ""), r.get("任務類型", ""))
            groups[key].append(r)

        row_i = 2
        for (name, cond, task), recs in groups.items():
            def _float(v):
                try: return float(v)
                except: return 0.0
            f_c_col = "穿戴壓迫力(N)" if cond == "exo" else "未穿戴壓迫力(N)"
            max_fc = max(_float(r.get(f_c_col, 0)) for r in recs)
            avg_red = sum(_float(r.get("減壓率(%)", 0)) for r in recs) / len(recs)
            high_risk = sum(1 for r in recs if r.get("風險等級") == "high")
            max_li = max(_float(r.get("負載指數", 0)) for r in recs)
            row_data = [name, cond, task, round(max_fc, 1), round(avg_red, 2), high_risk, round(max_li, 3)]
            for ci, val in enumerate(row_data, 1):
                cell = ws_sum.cell(row_i, ci, val)
                cell.font = Font(name="Arial")
                cell.alignment = Alignment(horizontal="center")
                cell.border = BORDER
            row_i += 1

        _auto_col_width(ws_sum)
        _save_wb(wb)
    return DB_PATH
